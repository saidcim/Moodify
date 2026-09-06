import os
import sys
import glob
import json
import time
import logging
import datetime
import threading
import openpyxl
import spotipy
from spotipy.oauth2 import SpotifyOAuth
from groq import Groq

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()],
)
log = logging.getLogger(__name__)


def _require_env(name: str) -> str:
    val = os.environ.get(name)
    if not val:
        raise SystemExit(
            f"ERROR: required environment variable is not set: {name}. "
            f"On GitHub Actions add it under repository Secrets; "
            f"for local runs see env.example."
        )
    return val


SPOTIFY_CLIENT_ID = _require_env("SPOTIFY_CLIENT_ID")
SPOTIFY_CLIENT_SECRET = _require_env("SPOTIFY_CLIENT_SECRET")
SPOTIFY_REDIRECT_URI = os.environ.get("SPOTIFY_REDIRECT_URI", "http://127.0.0.1:5000/callback")
SPOTIFY_MARKET = os.environ.get("SPOTIFY_MARKET", "US")
SPOTIFY_REFRESH_TOKEN = _require_env("SPOTIFY_REFRESH_TOKEN")
GROQ_API_KEY = _require_env("GROQ_API_KEY")
USER_NOTE = os.environ.get("USER_NOTE", "").strip()[:500]
PLAYLIST_NAME = os.environ.get("PLAYLIST_NAME", "AI Daily Mix")
PLAYLIST_SIZE = int(os.environ.get("PLAYLIST_SIZE", "40"))
DISCOVERY_QUOTA = min(int(os.environ.get("DISCOVERY_QUOTA", "8")), max(PLAYLIST_SIZE // 3, 1))

DATA_DIR = os.environ.get("DATA_DIR", "/data" if os.path.isdir("/data") else ".")
os.makedirs(DATA_DIR, exist_ok=True)
HISTORY_FILE_PREFIX = "playlist_history_cycle_"
STATE_FILE = os.path.join(DATA_DIR, "bot_state.json")
DASHBOARD_FILE = os.path.join(DATA_DIR, "dashboard_state.json")


def history_file_for_cycle(cycle: int) -> str:
    return os.path.join(DATA_DIR, f"{HISTORY_FILE_PREFIX}{int(cycle):03d}.xlsx")


def _latest_history_file() -> str | None:
    files = glob.glob(os.path.join(DATA_DIR, f"{HISTORY_FILE_PREFIX}*.xlsx"))
    if not files:
        return None

    def _cycle_num(p):
        try:
            return int(os.path.basename(p)[len(HISTORY_FILE_PREFIX):-len(".xlsx")])
        except ValueError:
            return -1

    return max(files, key=_cycle_num)


ARCHIVE_LIMIT = 50
USER_NOTE_WINDOW = 3
USER_NOTE_MAX_CHARS = 200

GROQ_MODEL_CHAIN = [m for m in [
    os.environ.get("GROQ_MODEL", ""),
    "openai/gpt-oss-120b",
    "llama-3.3-70b-versatile",
    "openai/gpt-oss-20b",
    "llama-3.1-8b-instant",
] if m]
_ACTIVE_GROQ_MODEL: str | None = None

GROQ_TPM_BUDGET = int(os.environ.get("GROQ_TPM_BUDGET", "7000"))
_GROQ_USAGE: list = []

DEFAULT_CARRY_OVER = 5
CARRY_OVER_MIN, CARRY_OVER_MAX = 0, 10
CARRY_OVER_FLOOR = 2
CARRY_TUNE_DEADBAND = 0.10

is_running = False
STATE_LOCK = threading.Lock()

AFFINITY_DECAY = 0.95
AFFINITY_PRUNE_BELOW = 0.05


def _default_user_profile() -> dict:
    return {
        "version": 1,
        "last_updated": None,
        "learned_patterns": [],
        "genre_affinity": {},
        "artist_affinity": {},
        "mood_music_map": {},
        "next_advice": "",
        "adaptation_metrics": {
            "cycles_completed": 0,
            "score_basis": "insufficient_engagement_history",
            "avg_score_first_3": 0.0,
            "avg_score_last_3": 0.0,
            "improvement_delta": 0.0,
            "avg_ai_score_last_3": 0.0,
            "avg_plays_trend": "0%",
            "carry_over_success_rate": None,
            "carry_over_measured": {},
            "confidence": 0.0,
        },
    }


def _default_dynamic_config() -> dict:
    return {
        "carry_over": DEFAULT_CARRY_OVER,
        "last_tuned_cycle": 0,
        "tune_reason": "",
    }


def _default_state() -> dict:
    return {
        "playlist_id": os.environ.get("PLAYLIST_ID", ""),
        "last_update": None,
        "cycle": 0,
        "feedback_history": [],
        "mood_history": [],
        "playlist_archive": [],
        "user_profile": _default_user_profile(),
        "dynamic_config": _default_dynamic_config(),
        "ai_notes": "",
        "cumulative_plays": {},
        "seen_play_events": [],
        "user_notes": [],
    }


def _load_state_from_file(path: str) -> dict | None:
    if not os.path.exists(path):
        return None
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        log.warning(f"Could not read state file ({path}): {e}")
        return None


def _merge_state_defaults(state: dict) -> dict:
    default = _default_state()
    for key, val in default.items():
        if key not in state:
            state[key] = val
    if "user_profile" in state:
        profile_default = _default_user_profile()
        for k, v in profile_default.items():
            if k not in state["user_profile"]:
                state["user_profile"][k] = v
        metrics = profile_default["adaptation_metrics"]
        for k, v in metrics.items():
            if k not in state["user_profile"].get("adaptation_metrics", {}):
                state["user_profile"]["adaptation_metrics"][k] = v

    dyn_default = _default_dynamic_config()
    dyn = state.get("dynamic_config") or {}
    dyn.pop("last_recovery_cycle", None)
    for k, v in dyn_default.items():
        if k not in dyn:
            dyn[k] = v
    dyn["carry_over"] = max(CARRY_OVER_MIN, min(CARRY_OVER_MAX, int(dyn["carry_over"])))
    dyn["carry_over"] = min(dyn["carry_over"], PLAYLIST_SIZE)
    state["dynamic_config"] = dyn
    return state


def load_state() -> dict:
    state = _load_state_from_file(STATE_FILE)
    if state is None and STATE_FILE != "bot_state.json":
        state = _load_state_from_file("bot_state.json")
    if state is None:
        log.info("No saved state found, starting clean.")
        state = _default_state()
    else:
        log.info(f"State loaded (cycle #{state.get('cycle', 0)}).")
    return _merge_state_defaults(state)


def write_dashboard_state(state: dict):
    try:
        payload = {
            "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "cycle": state.get("cycle", 0),
            "last_update": state.get("last_update"),
            "playlist_id": state.get("playlist_id"),
            "is_running": False,
            "feedback_history": state.get("feedback_history", [])[-30:],
            "mood_history": state.get("mood_history", [])[-30:],
            "user_profile": state.get("user_profile", {}),
            "dynamic_config": state.get("dynamic_config", {}),
            "cumulative_plays": state.get("cumulative_plays", {}),
            "user_notes": state.get("user_notes", [])[-10:],
            "playlist_archive_summary": state.get("playlist_archive", [])[-10:],
        }
        with open(DASHBOARD_FILE, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
    except Exception as e:
        log.warning(f"Could not write dashboard_state.json: {e}")


def save_state(state: dict):
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, indent=2, ensure_ascii=False)
    write_dashboard_state(state)


TRACK_HEADERS = ["Cycle", "Date", "Time", "Spotify ID", "Track", "Artist", "Album",
                 "AI Score", "Play Count"]
SUMMARY_HEADERS = ["Cycle", "Datetime", "Mood", "Energy", "Score", "Track Count",
                   "Avg Plays", "Carried Tracks", "AI Analysis Summary"]


def _get_or_create_sheets(wb):
    if "Tracks" in wb.sheetnames:
        ws_tracks = wb["Tracks"]
    else:
        ws_tracks = wb.active
        ws_tracks.title = "Tracks"
    if "Cycle Summary" in wb.sheetnames:
        ws_summary = wb["Cycle Summary"]
    else:
        ws_summary = wb.create_sheet("Cycle Summary")
    if ws_tracks.max_row == 1 and ws_tracks.cell(1, 1).value is None:
        for col, h in enumerate(TRACK_HEADERS, 1):
            ws_tracks.cell(1, col).value = h
    if ws_summary.max_row == 1 and ws_summary.cell(1, 1).value is None:
        for col, h in enumerate(SUMMARY_HEADERS, 1):
            ws_summary.cell(1, col).value = h
    return ws_tracks, ws_summary


def save_to_excel(tracks: list, cycle: int, score=None, play_counts: dict | None = None,
                  history_file: str | None = None):
    play_counts = play_counts or {}
    path = history_file or history_file_for_cycle(cycle)
    now = datetime.datetime.now()
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M:%S")

    wb = openpyxl.load_workbook(path) if os.path.exists(path) else openpyxl.Workbook()
    ws_tracks, _ = _get_or_create_sheets(wb)

    for track in tracks:
        plays = play_counts.get(track["id"], 0)
        ws_tracks.append([
            cycle, date_str, time_str, track["id"], track["name"],
            track["artist"], track["album"], score if score is not None else "-", plays,
        ])
    wb.save(path)
    log.info(f"Saved {len(tracks)} tracks to Excel (cycle {cycle}) -> {os.path.basename(path)}")


def save_cycle_summary_to_excel(cycle: int, mood_data: dict, score: float, track_count: int,
                                avg_plays: float, carry_count: int, analysis: str,
                                history_file: str | None = None):
    path = history_file or history_file_for_cycle(cycle)
    now = datetime.datetime.now()
    datetime_str = now.strftime("%Y-%m-%d %H:%M:%S")

    wb = openpyxl.load_workbook(path) if os.path.exists(path) else openpyxl.Workbook()
    _, ws_summary = _get_or_create_sheets(wb)
    ws_summary.append([
        cycle, datetime_str,
        mood_data.get("mood", "-"),
        mood_data.get("energy_level", "-"),
        score, track_count, round(avg_plays, 2), carry_count,
        (analysis or "")[:200],
    ])
    wb.save(path)
    log.info(f"Saved cycle summary to Excel (cycle {cycle}) -> {os.path.basename(path)}")


def get_spotify() -> spotipy.Spotify:
    auth = SpotifyOAuth(
        client_id=SPOTIFY_CLIENT_ID,
        client_secret=SPOTIFY_CLIENT_SECRET,
        redirect_uri=SPOTIFY_REDIRECT_URI,
        scope=(
            "user-read-recently-played user-top-read user-library-read "
            "playlist-modify-public playlist-modify-private playlist-read-private"
        ),
    )
    token_info = auth.refresh_access_token(SPOTIFY_REFRESH_TOKEN)
    return spotipy.Spotify(auth=token_info["access_token"])


def _parse_played_at(raw: str) -> datetime.datetime:
    dt = datetime.datetime.fromisoformat(raw.replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=datetime.timezone.utc)
    return dt


def get_listening_data(sp: spotipy.Spotify) -> dict:
    data = {}
    recent = sp.current_user_recently_played(limit=50)
    recent_tracks = []
    cutoff = datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(days=3)
    for item in recent["items"]:
        played_at = _parse_played_at(item["played_at"])
        if played_at >= cutoff:
            t = item["track"]
            recent_tracks.append({
                "id": t["id"], "name": t["name"],
                "artist": t["artists"][0]["name"], "album": t["album"]["name"],
                "played_at": item["played_at"],
            })
    data["recent_tracks"] = recent_tracks

    top_short = sp.current_user_top_tracks(limit=50, time_range="short_term")
    data["top_short"] = [{"id": t["id"], "name": t["name"], "artist": t["artists"][0]["name"],
                          "popularity": t["popularity"]} for t in top_short["items"]]

    top_medium = sp.current_user_top_tracks(limit=50, time_range="medium_term")
    data["top_medium"] = [{"id": t["id"], "name": t["name"], "artist": t["artists"][0]["name"],
                           "popularity": t["popularity"]} for t in top_medium["items"]]

    top_artists = sp.current_user_top_artists(limit=20, time_range="short_term")
    data["top_artists"] = [{"id": a["id"], "name": a["name"], "genres": a["genres"],
                            "popularity": a.get("popularity", 0)} for a in top_artists["items"]]

    saved = sp.current_user_saved_tracks(limit=50)
    data["saved_tracks"] = [{"id": item["track"]["id"], "name": item["track"]["name"],
                             "artist": item["track"]["artists"][0]["name"]} for item in saved["items"]]

    return data


def get_playlist_play_counts(playlist_track_ids, recent_tracks):
    counts = {tid: 0 for tid in playlist_track_ids}
    for t in recent_tracks:
        if t["id"] in counts:
            counts[t["id"]] += 1
    return counts


def build_listening_fingerprint(listening_data: dict) -> dict:
    genres = []
    for a in listening_data.get("top_artists", [])[:10]:
        genres.extend(a.get("genres", [])[:3])
    genre_counts: dict[str, int] = {}
    for g in genres:
        genre_counts[g] = genre_counts.get(g, 0) + 1
    top_genres = sorted(genre_counts, key=genre_counts.get, reverse=True)[:5]
    return {
        "top_artists": [a["name"] for a in listening_data.get("top_artists", [])[:8]],
        "top_genres": top_genres,
        "track_count_3d": len(listening_data.get("recent_tracks", [])),
    }


def find_or_create_playlist(sp: spotipy.Spotify, state: dict) -> str:
    user_id = sp.current_user()["id"]
    if state.get("playlist_id"):
        try:
            pl = sp.playlist(state["playlist_id"])
            log.info(f"Existing playlist: {pl['name']} ({pl['id']})")
            return state["playlist_id"]
        except Exception:
            log.warning("Playlist ID is invalid, searching...")

    matching = []
    offset = 0
    while True:
        results = sp.current_user_playlists(limit=50, offset=offset)
        for pl in results["items"]:
            if pl["name"] == PLAYLIST_NAME and pl["owner"]["id"] == user_id:
                matching.append(pl)
        if results["next"] is None:
            break
        offset += 50

    if matching:
        keeper = matching[0]
        for dup in matching[1:]:
            try:
                sp.current_user_unfollow_playlist(dup["id"])
                log.info(f"Duplicate removed: {dup['id']}")
            except Exception as e:
                log.warning(f"Could not remove duplicate: {e}")
        return keeper["id"]

    pl = sp.user_playlist_create(
        user=user_id, name=PLAYLIST_NAME, public=False,
        description="AI playlist refreshed every 2 days",
    )
    log.info(f"New playlist: {pl['id']}")
    return pl["id"]


def update_playlist(sp, playlist_id, track_ids):
    sp.playlist_replace_items(playlist_id, [])
    for i in range(0, len(track_ids), 100):
        chunk = [f"spotify:track:{tid}" for tid in track_ids[i:i + 100]]
        sp.playlist_add_items(playlist_id, chunk)
    log.info(f"Playlist updated: {len(track_ids)} tracks")


def build_discovery_candidates(sp, listening_data, known_ids: set, state: dict) -> list:
    discoveries = []
    seen = set(known_ids)
    for a in state.get("playlist_archive", [])[-6:]:
        seen.update(a.get("discovery_ids", []))

    top_artists = listening_data.get("top_artists", [])

    for artist in top_artists[:8]:
        if not artist.get("id"):
            continue
        try:
            tt = sp.artist_top_tracks(artist["id"], country=SPOTIFY_MARKET)
            added = 0
            for t in tt.get("tracks", []):
                if t.get("id") and t["id"] not in seen:
                    discoveries.append({"id": t["id"], "name": t["name"],
                                        "artist": t["artists"][0]["name"],
                                        "source": "known artist, unplayed track"})
                    seen.add(t["id"])
                    added += 1
                if added >= 2:
                    break
        except Exception as e:
            log.warning(f"artist_top_tracks failed ({artist.get('name')}): {e}")

    fingerprint_genres = build_listening_fingerprint(listening_data)["top_genres"]
    known_artist_names = {a["name"].lower() for a in top_artists}
    for genre in fingerprint_genres[:3]:
        try:
            res = sp.search(q=f'genre:"{genre}"', type="artist", limit=10)
            for a in res.get("artists", {}).get("items", []):
                pop = a.get("popularity", 0)
                if a["name"].lower() in known_artist_names or not (30 <= pop <= 70):
                    continue
                try:
                    tt = sp.artist_top_tracks(a["id"], country=SPOTIFY_MARKET)
                except Exception:
                    continue
                added = 0
                for t in tt.get("tracks", []):
                    if t.get("id") and t["id"] not in seen:
                        discoveries.append({"id": t["id"], "name": t["name"],
                                            "artist": a["name"],
                                            "source": f"genre discovery: {genre}"})
                        seen.add(t["id"])
                        added += 1
                    if added >= 2:
                        break
                known_artist_names.add(a["name"].lower())
                if added:
                    break
        except Exception as e:
            log.warning(f"Genre search failed ({genre}): {e}")

    log.info(f"Discovery pool: {len(discoveries)} candidates")
    return discoveries[:DISCOVERY_QUOTA * 3]


def measure_carry_over_performance(state: dict, play_counts: dict,
                                   current_track_ids: list) -> dict:
    archive = state.get("playlist_archive", [])
    prev_carry = (archive[-1].get("carry_ids") or []) if archive else []
    carry_ids = [tid for tid in prev_carry if tid in play_counts]
    carry_set = set(carry_ids)
    fresh_ids = [tid for tid in current_track_ids if tid not in carry_set]

    def _played_ratio(ids):
        if not ids:
            return None
        return sum(1 for tid in ids if play_counts.get(tid, 0) > 0) / len(ids)

    return {
        "carry_played_ratio": _played_ratio(carry_ids),
        "fresh_played_ratio": _played_ratio(fresh_ids),
        "carry_sample": len(carry_ids),
        "fresh_sample": len(fresh_ids),
    }


def analyze_patterns(state: dict, listening_data: dict, play_counts: dict,
                     current_track_ids: list) -> dict:
    profile = state.setdefault("user_profile", _default_user_profile())
    archive = state.get("playlist_archive", [])
    feedback = state.get("feedback_history", [])
    fingerprint = build_listening_fingerprint(listening_data)

    genre_affinity: dict[str, float] = dict(profile.get("genre_affinity", {}))
    artist_affinity: dict[str, float] = dict(profile.get("artist_affinity", {}))
    mood_music_map: dict[str, list] = dict(profile.get("mood_music_map", {}))

    genre_affinity = {k: round(v * AFFINITY_DECAY, 3) for k, v in genre_affinity.items()
                      if v * AFFINITY_DECAY >= AFFINITY_PRUNE_BELOW}
    artist_affinity = {k: round(v * AFFINITY_DECAY, 3) for k, v in artist_affinity.items()
                       if v * AFFINITY_DECAY >= AFFINITY_PRUNE_BELOW}

    for genre in fingerprint["top_genres"]:
        genre_affinity[genre] = min(1.0, genre_affinity.get(genre, 0.0) + 0.1)
    for artist in fingerprint["top_artists"][:5]:
        artist_affinity[artist] = min(1.0, artist_affinity.get(artist, 0.0) + 0.1)

    if archive and feedback:
        last_archive = archive[-1]
        last_mood = last_archive.get("mood", "")
        if last_mood and fingerprint["top_genres"]:
            existing = mood_music_map.get(last_mood, [])
            for g in fingerprint["top_genres"][:3]:
                if g not in existing:
                    existing.append(g)
            mood_music_map[last_mood] = existing[:5]

    carry_perf = measure_carry_over_performance(state, play_counts, current_track_ids)
    carry_rate = carry_perf["carry_played_ratio"]

    eng_scores = [f["engagement_score"] for f in feedback
                  if f.get("engagement_score") is not None]
    ai_scores = [f["score"] for f in feedback if f.get("score") is not None]

    if len(eng_scores) >= 2:
        series, basis = eng_scores, "engagement"
    else:
        series, basis = [], "insufficient_engagement_history"

    avg_first_3 = sum(series[:3]) / len(series[:3]) if series else 0.0
    avg_last_3 = sum(series[-3:]) / len(series[-3:]) if series else 0.0
    improvement = round(avg_last_3 - avg_first_3, 2) if len(series) >= 2 else 0.0

    plays_trend = "0%"
    if len(archive) >= 2:
        early = sum(a.get("avg_plays", 0) for a in archive[:3]) / min(3, len(archive))
        recent = sum(a.get("avg_plays", 0) for a in archive[-3:]) / min(3, len(archive))
        if early > 0:
            pct = int((recent - early) / early * 100)
            plays_trend = f"{pct:+d}%"

    profile["genre_affinity"] = genre_affinity
    profile["artist_affinity"] = artist_affinity
    profile["mood_music_map"] = mood_music_map
    profile["adaptation_metrics"] = {
        "cycles_completed": state.get("cycle", 0),
        "score_basis": basis,
        "avg_score_first_3": round(avg_first_3, 2),
        "avg_score_last_3": round(avg_last_3, 2),
        "improvement_delta": improvement,
        "avg_ai_score_last_3": round(sum(ai_scores[-3:]) / len(ai_scores[-3:]), 2) if ai_scores else 0.0,
        "avg_plays_trend": plays_trend,
        "carry_over_success_rate": round(carry_rate, 2) if carry_rate is not None else None,
        "carry_over_measured": carry_perf,
        "confidence": profile.get("adaptation_metrics", {}).get("confidence", 0.0),
    }
    profile["last_updated"] = datetime.datetime.now(datetime.timezone.utc).isoformat()
    return profile


def tune_dynamic_config(state: dict, play_counts: dict, current_track_ids: list) -> dict:
    dyn = state.get("dynamic_config", _default_dynamic_config())
    cycle_now = state.get("cycle", 0)
    before = int(dyn.get("carry_over", DEFAULT_CARRY_OVER))

    if before < CARRY_OVER_FLOOR:
        dyn["carry_over"] = CARRY_OVER_FLOOR
        dyn["last_tuned_cycle"] = cycle_now
        dyn["tune_reason"] = (f"Floor correction: carry_over {before} -> {CARRY_OVER_FLOOR} "
                              f"(a playlist cannot be 100% unfamiliar tracks).")
        log.info(dyn["tune_reason"])
        return dyn

    perf = measure_carry_over_performance(state, play_counts, current_track_ids)
    carry_r, fresh_r = perf["carry_played_ratio"], perf["fresh_played_ratio"]

    if carry_r is None or fresh_r is None or perf["carry_sample"] < 2:
        dyn["tune_reason"] = (f"Not enough carry-over data (n={perf['carry_sample']}), "
                              f"keeping carry_over at {before}.")
        log.info(dyn["tune_reason"])
        return dyn

    diff = carry_r - fresh_r
    if diff > CARRY_TUNE_DEADBAND:
        new_carry, why = min(before + 1, DEFAULT_CARRY_OVER), "carried tracks played more than fresh ones"
    elif diff < -CARRY_TUNE_DEADBAND:
        new_carry, why = max(before - 1, CARRY_OVER_FLOOR), "carried tracks played less than fresh ones"
    else:
        new_carry, why = before, "difference inside the deadband"

    dyn["carry_over"] = new_carry
    dyn["last_tuned_cycle"] = cycle_now
    dyn["tune_reason"] = (
        f"{why}: carried {carry_r * 100:.0f}% / fresh {fresh_r * 100:.0f}% "
        f"(n={perf['carry_sample']}/{perf['fresh_sample']}) -> carry_over {before}->{new_carry}"
    )
    log.info(f"Dynamic config: {dyn['tune_reason']}")
    return dyn


def _parse_ai_json(raw: str) -> dict:
    raw = (raw or "").strip()
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    raw = raw.strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start = raw.find("{")
        if start == -1:
            raise
        depth, in_str, esc = 0, False, False
        for i, ch in enumerate(raw[start:], start):
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    in_str = False
            elif ch == '"':
                in_str = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return json.loads(raw[start:i + 1])
        raise


def _compact_tracks(tracks: list, limit: int) -> list:
    return [f'{t.get("name", "?")} - {t.get("artist", "?")}' for t in tracks[:limit]]


def _compact_artists(artists: list, limit: int) -> list:
    return [{"name": a.get("name", "?"), "genres": (a.get("genres") or [])[:2]}
            for a in artists[:limit]]


def _compact_archive(archive: list) -> list:
    return [{"cycle": a.get("cycle"), "mood": a.get("mood"),
             "engagement": a.get("engagement_score"), "avg_plays": a.get("avg_plays"),
             "carry": a.get("carry_over_count"), "discovery": a.get("discovery_count")}
            for a in archive]


def _model_extra_body(model: str) -> dict:
    if model.startswith("openai/gpt-oss"):
        return {"reasoning_effort": "low", "include_reasoning": False}
    return {}


def _estimate_tokens(text: str) -> int:
    return len(text) // 3 + 200


def _pace_for_tokens(need: int):
    if need >= GROQ_TPM_BUDGET:
        log.warning(f"Single request exceeds the TPM budget (~{need} > {GROQ_TPM_BUDGET}); trying anyway.")
        return
    while True:
        now = time.time()
        _GROQ_USAGE[:] = [(t, n) for t, n in _GROQ_USAGE if now - t < 60]
        used = sum(n for _, n in _GROQ_USAGE)
        if not _GROQ_USAGE or used + need <= GROQ_TPM_BUDGET:
            return
        wait = max(61 - (now - _GROQ_USAGE[0][0]), 1)
        log.info(f"Groq TPM window is full (~{used}+{need} > {GROQ_TPM_BUDGET}), waiting {wait:.0f}s.")
        time.sleep(wait)


def groq_json(client: Groq, prompt: str, max_tokens: int) -> dict:
    global _ACTIVE_GROQ_MODEL
    chain = ([_ACTIVE_GROQ_MODEL] if _ACTIVE_GROQ_MODEL else []) + \
            [m for m in GROQ_MODEL_CHAIN if m != _ACTIVE_GROQ_MODEL]
    est = _estimate_tokens(prompt) + max_tokens
    last_err = None

    for model in chain:
        extra = _model_extra_body(model)
        for force_json in (True, False):
            kwargs = {"response_format": {"type": "json_object"}} if force_json else {}
            if extra:
                kwargs["extra_body"] = extra
            try:
                _pace_for_tokens(est)
                resp = client.chat.completions.create(
                    model=model, max_tokens=max_tokens,
                    messages=[{"role": "user", "content": prompt}], **kwargs,
                )
                _GROQ_USAGE.append((time.time(), est))
                choice = resp.choices[0]
                content = (choice.message.content or "").strip()
                if not content:
                    reason = getattr(choice, "finish_reason", "?")
                    log.warning(f"{model}: empty response (finish_reason={reason}), moving to the next attempt.")
                    last_err = RuntimeError(f"{model} returned an empty response (finish_reason={reason})")
                    continue
                if model != _ACTIVE_GROQ_MODEL:
                    log.info(f"Groq model: {model}")
                    _ACTIVE_GROQ_MODEL = model
                return _parse_ai_json(content)
            except Exception as e:
                msg = str(e)
                _GROQ_USAGE.append((time.time(), est))
                if force_json and ("json_validate_failed" in msg or "Failed to validate JSON" in msg):
                    log.warning(f"{model}: json_object mode failed, retrying without format enforcement.")
                    continue
                if "model_not_found" in msg or "does not exist" in msg or "decommissioned" in msg:
                    log.warning(f"Groq model unavailable ({model}), moving to the next one.")
                    last_err = e
                    break
                raise
        else:
            if _ACTIVE_GROQ_MODEL == model:
                _ACTIVE_GROQ_MODEL = None
            log.warning(f"{model} produced no usable response, moving to the next model.")
    raise RuntimeError(
        f"No Groq model was usable. Tried: {chain}. Last error: {last_err}"
    )


def synthesize_user_profile(state: dict, listening_data: dict, client: Groq) -> dict:
    profile = state.get("user_profile", _default_user_profile())
    archive = _compact_archive(state.get("playlist_archive", [])[-4:])
    fingerprint = build_listening_fingerprint(listening_data)

    if not archive:
        return profile

    prompt = f"""You are a music taste analyst. Study the listening history and extract learned patterns.

## Current Profile
{json.dumps(profile.get("adaptation_metrics", {}), ensure_ascii=False)}

## Recent Cycles
{json.dumps(archive, ensure_ascii=False)}

## Listening Summary For This Period
{json.dumps(fingerprint, ensure_ascii=False)}

## Existing Learned Patterns
{json.dumps(profile.get("learned_patterns", [])[-5:], ensure_ascii=False)}

Task: discover new patterns that describe this listener better. Keep the existing patterns and add the new ones (8 items max in total).

JSON ONLY:
{{"learned_patterns": ["pattern1", "pattern2"], "next_cycle_advice": "advice for the next cycle", "confidence": 0.72}}"""

    try:
        parsed = groq_json(client, prompt, max_tokens=500)
        existing = profile.get("learned_patterns", [])
        new_patterns = parsed.get("learned_patterns", [])
        merged = list(existing)
        for p in new_patterns:
            if p not in merged:
                merged.append(p)
        profile["learned_patterns"] = merged[-8:]
        profile["next_advice"] = parsed.get("next_cycle_advice", profile.get("next_advice", ""))
        profile["adaptation_metrics"]["confidence"] = float(parsed.get("confidence", 0.0))
        log.info(f"Profile synthesis: {len(profile['learned_patterns'])} patterns, "
                 f"confidence {profile['adaptation_metrics']['confidence'] * 100:.0f}%")
    except Exception as e:
        log.warning(f"Profile synthesis failed: {e}")

    return profile


def analyze_mood(listening_data, client):
    all_recent = listening_data.get("recent_tracks", [])
    prompt = f"""You are a music psychology expert. Analyse the listener's mood from the last 3 days of listening data.

## Tracks Played In The Last 3 Days
{json.dumps(_compact_tracks(all_recent, 20), ensure_ascii=False)}

## Short Term Top Tracks
{json.dumps(_compact_tracks(listening_data.get('top_short', []), 10), ensure_ascii=False)}

## Favourite Artists
{json.dumps(_compact_artists(listening_data.get('top_artists', []), 6), ensure_ascii=False)}

Return JSON ONLY:
{{"mood": "main mood in English", "mood_emoji": "emoji", "energy_level": "low/medium/high",
"dominant_genres": ["genre1","genre2"], "top_artists_this_period": ["artist1","artist2","artist3"],
"summary": "2-3 sentence summary in English", "track_count": {len(all_recent)}}}"""

    try:
        return groq_json(client, prompt, max_tokens=600)
    except Exception as e:
        log.warning(f"Mood analysis failed: {e}")
        return {"mood": "Unknown", "mood_emoji": "?", "energy_level": "medium",
                "dominant_genres": [], "top_artists_this_period": [],
                "summary": "Analysis unavailable.", "track_count": len(all_recent)}


def _build_user_notes_block(state: dict) -> str:
    notes = state.get("user_notes", [])[-USER_NOTE_WINDOW:]
    if not notes:
        return ""
    lines = [f'- (cycle #{n.get("cycle", "?")}) {str(n.get("note", ""))[:USER_NOTE_MAX_CHARS]}'
             for n in notes]
    return ("\n## DIRECT USER REQUESTS - HIGHEST PRIORITY SIGNAL\n"
            "These were written by the user. They outrank every other signal\n"
            "(mood, history, profile); the last one is the most recent.\n"
            + "\n".join(lines) + "\n")


def _build_learning_context(state: dict) -> str:
    profile = state.get("user_profile", {})
    metrics = profile.get("adaptation_metrics", {})
    patterns = profile.get("learned_patterns", [])[-5:]
    advice = profile.get("next_advice", "")
    archive_summary = _compact_archive(state.get("playlist_archive", [])[-3:])

    def _top_affinity(d, n):
        return dict(sorted(d.items(), key=lambda kv: kv[1], reverse=True)[:n])

    lines = [
        f"## User Profile (learned, cycle #{state.get('cycle', 0)})",
        f"Genre affinity: {json.dumps(_top_affinity(profile.get('genre_affinity', {}), 8), ensure_ascii=False)}",
        f"Artist affinity: {json.dumps(_top_affinity(profile.get('artist_affinity', {}), 6), ensure_ascii=False)}",
        "",
        "## Learned Patterns",
        "\n".join(f"- {p}" for p in patterns) if patterns else "- No patterns yet",
        "",
        "## Past Cycle Comparison",
        json.dumps(archive_summary, ensure_ascii=False),
        "",
        "## Adaptation Note",
        (f"Confidence: {metrics.get('confidence', 0) * 100:.0f}% | "
         f"Improvement (real listening signal): {metrics.get('improvement_delta', 0):+.1f} | "
         f"Play trend: {metrics.get('avg_plays_trend', '0%')}"
         if metrics.get("score_basis") == "engagement"
         else f"Confidence: {metrics.get('confidence', 0) * 100:.0f}% | "
              f"Improvement: not measurable yet (not enough engagement history) | "
              f"Play trend: {metrics.get('avg_plays_trend', '0%')}"),
        f"Advice: {advice}" if advice else "",
    ]
    return "\n".join(lines)


def ai_analyze_and_build(listening_data, play_counts, state, candidates,
                         discovery_candidates, current_track_ids=None):
    client = Groq(api_key=GROQ_API_KEY)
    current_track_ids = current_track_ids or []

    dyn = state.get("dynamic_config", _default_dynamic_config())
    carry_over_limit = dyn["carry_over"]
    playlist_size = PLAYLIST_SIZE

    meta = {c["id"]: c for c in candidates if c.get("id")}
    disc_meta = {d["id"]: d for d in discovery_candidates if d.get("id")}
    disc_ids = list(disc_meta.keys())

    if current_track_ids and play_counts:
        sorted_old = sorted(play_counts.items(), key=lambda x: x[1], reverse=True)
        carry_over_pool = [tid for tid, count in sorted_old if count > 0][:carry_over_limit]
        banned_ids = [tid for tid in current_track_ids if tid not in carry_over_pool]
    else:
        carry_over_pool = []
        banned_ids = list(current_track_ids)

    banned_set = set(banned_ids)
    fresh_ids = [tid for tid in meta if tid not in banned_set and tid not in disc_meta]
    log.info(f"Candidates: {len(fresh_ids)} fresh + {len(carry_over_pool)} carryable + {len(disc_ids)} discovery")

    def _fmt(tid, extra=""):
        m = meta.get(tid) or disc_meta.get(tid) or {}
        label = f'{m.get("name", "?")} - {m.get("artist", "?")}'
        return {"id": tid, "track": label, **({"info": extra} if extra else {})}

    fresh_view = [_fmt(t) for t in fresh_ids[:35]]
    carry_view = [{**_fmt(t), "plays": play_counts.get(t, 0)} for t in carry_over_pool]
    disc_view = [{"id": d["id"], "track": f'{d["name"]} - {d["artist"]}', "source": d["source"]}
                 for d in discovery_candidates[:DISCOVERY_QUOTA * 3]]

    learning_context = _build_learning_context(state)
    mood_data = analyze_mood(listening_data, client)
    log.info(f"Mood: {mood_data.get('mood')} {mood_data.get('mood_emoji')}")

    user_notes_block = _build_user_notes_block(state)

    prompt = f"""You are a music curator. Pick a playlist of {playlist_size} tracks.
{user_notes_block}
{learning_context}

Cycle #{state['cycle'] + 1}
Mood: {json.dumps(mood_data, ensure_ascii=False)}
Recently played: {json.dumps(_compact_tracks(listening_data['recent_tracks'], 12), ensure_ascii=False)}
Top tracks: {json.dumps(_compact_tracks(listening_data['top_short'], 10), ensure_ascii=False)}
Top artists: {json.dumps(_compact_artists(listening_data['top_artists'], 6), ensure_ascii=False)}

## FRESH CANDIDATE TRACKS (build the core from these)
{json.dumps(fresh_view, ensure_ascii=False)}

## DISCOVERY CANDIDATES (tracks the user has never played - pick EXACTLY {DISCOVERY_QUOTA}, the ones that best fit the mood and profile)
{json.dumps(disc_view, ensure_ascii=False)}

## CARRYABLE OLD TRACKS (played a lot last period, you may keep at most {carry_over_limit})
{json.dumps(carry_view, ensure_ascii=False)}

Task:
0. If a "DIRECT USER REQUESTS" section appears above, follow it FIRST; when the most
   recent request conflicts with any other signal, the request wins. State in one
   sentence how you satisfied it inside the analysis text.
1. Choose according to the learned profile and patterns - show that you know this listener.
2. Add EXACTLY {DISCOVERY_QUOTA} tracks from the discovery candidates (discovery quota).
3. You may add at most {carry_over_limit} tracks from the carryable old tracks.
4. Fill the remaining slots from the fresh candidates. Use ONLY the ids listed above.
5. Match the mood, add variety, shuffle the order.

JSON ONLY:
{{"track_ids":["id1","id2"],"score":7.5,"analysis":"analysis in English","notes":"note for the next cycle"}}"""

    parsed = groq_json(client, prompt, max_tokens=2000)

    allowed = set(fresh_ids) | set(carry_over_pool) | set(disc_ids)
    track_ids, seen = [], set()
    for tid in parsed.get("track_ids", []):
        if tid in allowed and tid not in seen:
            track_ids.append(tid)
            seen.add(tid)

    carry_used = [tid for tid in track_ids if tid in carry_over_pool]
    if len(carry_used) > carry_over_limit:
        excess = set(carry_used[carry_over_limit:])
        track_ids = [tid for tid in track_ids if tid not in excess]
        log.info(f"Carry-over limit: removed {len(excess)} excess tracks")

    disc_used = [tid for tid in track_ids if tid in disc_meta]
    if len(disc_used) < DISCOVERY_QUOTA:
        for tid in disc_ids:
            if len(disc_used) >= DISCOVERY_QUOTA or len(track_ids) >= playlist_size:
                break
            if tid not in seen:
                track_ids.append(tid)
                seen.add(tid)
                disc_used.append(tid)

    for tid in fresh_ids:
        if len(track_ids) >= playlist_size:
            break
        if tid not in seen:
            track_ids.append(tid)
            seen.add(tid)
    track_ids = track_ids[:playlist_size]

    carry_used_ids = [tid for tid in track_ids if tid in carry_over_pool]
    discovery_used_ids = [tid for tid in track_ids if tid in disc_meta]
    fresh_count = len(track_ids) - len(carry_used_ids) - len(discovery_used_ids)
    log.info(f"AI finished. Score: {parsed.get('score')}, Fresh: {fresh_count}, "
             f"Discovery: {len(discovery_used_ids)}, Carried: {len(carry_used_ids)}, "
             f"Total: {len(track_ids)}")
    return (track_ids, parsed.get("notes", ""), float(parsed.get("score", 5.0)),
            parsed.get("analysis", ""), mood_data, carry_used_ids, discovery_used_ids)


def run_cycle(manual=False):
    global is_running
    if is_running:
        log.warning("Already running.")
        return {"status": "already_running"}
    is_running = True
    STATE_LOCK.acquire()
    trigger = "Manual" if manual else "Automatic"
    log.info(f"=========== Cycle Starting ({trigger}) ===========")

    try:
        state = load_state()
        sp = get_spotify()

        playlist_id = find_or_create_playlist(sp, state)
        state["playlist_id"] = playlist_id
        save_state(state)

        current_tracks = []
        if state.get("last_update"):
            try:
                items = sp.playlist_items(playlist_id, fields="items(track(id,name,artists,album))")
                for item in items["items"]:
                    t = item["track"]
                    if t and t.get("id"):
                        current_tracks.append({
                            "id": t["id"], "name": t["name"],
                            "artist": t["artists"][0]["name"], "album": t["album"]["name"],
                        })
            except Exception as e:
                log.warning(f"Could not read playlist: {e}")

        listening_data = get_listening_data(sp)
        fingerprint = build_listening_fingerprint(listening_data)

        play_counts = {}
        avg_plays = 0.0
        current_ids_for_ai = [t["id"] for t in current_tracks]
        if current_tracks:
            window_counts = get_playlist_play_counts(current_ids_for_ai, listening_data["recent_tracks"])
            cumulative = state.get("cumulative_plays", {})
            play_counts = {tid: max(window_counts.get(tid, 0), cumulative.get(tid, 0))
                           for tid in current_ids_for_ai}
            avg_plays = sum(play_counts.values()) / max(len(play_counts), 1)

        if USER_NOTE:
            state.setdefault("user_notes", []).append({
                "cycle": state.get("cycle", 0) + 1,
                "date": datetime.datetime.now(datetime.timezone.utc).isoformat(),
                "note": USER_NOTE,
            })
            state["user_notes"] = state["user_notes"][-10:]
            log.info(f"User request received: {USER_NOTE[:80]}")

        analyze_patterns(state, listening_data, play_counts, current_ids_for_ai)

        groq_client = Groq(api_key=GROQ_API_KEY)
        state["user_profile"] = synthesize_user_profile(state, listening_data, groq_client)
        state["dynamic_config"] = tune_dynamic_config(state, play_counts, current_ids_for_ai)

        candidates_map = {}
        for t in (listening_data["top_short"] + listening_data["top_medium"] +
                  listening_data["saved_tracks"] + listening_data["recent_tracks"]):
            if t.get("id") and t["id"] not in candidates_map:
                candidates_map[t["id"]] = {"id": t["id"], "name": t.get("name", "?"),
                                           "artist": t.get("artist", "?")}
        candidates = list(candidates_map.values())

        known_ids = set(candidates_map) | set(current_ids_for_ai)
        try:
            discovery_candidates = build_discovery_candidates(sp, listening_data, known_ids, state)
        except Exception as e:
            log.warning(f"Could not build discovery candidates: {e}")
            discovery_candidates = []

        (new_track_ids, ai_notes, score, analysis, mood_data,
         carry_used_ids, discovery_used_ids) = ai_analyze_and_build(
            listening_data, play_counts, state, candidates, discovery_candidates, current_ids_for_ai,
        )
        carry_count = len(carry_used_ids)

        engagement_score = None
        if current_tracks:
            played_ratio = len([c for c in play_counts.values() if c > 0]) / max(len(play_counts), 1)
            engagement_score = round(10 * (0.6 * played_ratio + 0.4 * min(avg_plays / 3.0, 1.0)), 1)

        completed_cycle = state["cycle"] + 1
        cycle_file = history_file_for_cycle(completed_cycle)

        archiving_cycle = state["cycle"]
        old_score = state["feedback_history"][-1]["score"] if state.get("feedback_history") else None
        if current_tracks:
            save_to_excel(current_tracks, archiving_cycle, old_score, play_counts, history_file=cycle_file)

        update_playlist(sp, playlist_id, new_track_ids)

        now = datetime.datetime.now(datetime.timezone.utc).isoformat()
        state["cycle"] += 1
        state["last_update"] = now
        state["ai_notes"] = ai_notes

        top_played = sorted(
            [{"id": tid, "name": next((t["name"] for t in current_tracks if t["id"] == tid), tid),
              "plays": c} for tid, c in play_counts.items() if c > 0],
            key=lambda x: x["plays"], reverse=True,
        )[:5]

        archive_entry = {
            "cycle": state["cycle"],
            "archived_at": now,
            "score": score,
            "avg_plays": round(avg_plays, 2),
            "mood": mood_data.get("mood", ""),
            "energy": mood_data.get("energy_level", ""),
            "track_count": len(new_track_ids),
            "top_played": top_played,
            "top_artists": fingerprint["top_artists"][:5],
            "genres": mood_data.get("dominant_genres", fingerprint["top_genres"][:5]),
            "carry_over_count": carry_count,
            "fresh_count": len(new_track_ids) - carry_count - len(discovery_used_ids),
            "discovery_count": len(discovery_used_ids),
            "discovery_ids": discovery_used_ids,
            "carry_ids": carry_used_ids,
            "engagement_score": engagement_score,
            "listening_fingerprint": fingerprint,
        }
        if "playlist_archive" not in state:
            state["playlist_archive"] = []
        state["playlist_archive"].append(archive_entry)
        state["playlist_archive"] = state["playlist_archive"][-ARCHIVE_LIMIT:]

        save_cycle_summary_to_excel(
            state["cycle"], mood_data, score, len(new_track_ids),
            avg_plays, carry_count, analysis, history_file=cycle_file,
        )

        if "mood_history" not in state:
            state["mood_history"] = []
        state["mood_history"].append({"date": now, "cycle": state["cycle"],
                                      "trigger": trigger, **mood_data})

        if "feedback_history" not in state:
            state["feedback_history"] = []
        state["feedback_history"].append({
            "cycle": state["cycle"], "date": now,
            "score": score, "engagement_score": engagement_score,
            "avg_plays": avg_plays, "analysis": analysis, "notes": ai_notes,
        })

        state["cumulative_plays"] = {}

        save_state(state)
        log.info(f"=========== Cycle #{state['cycle']} Completed ===========")
        return {"status": "ok", "cycle": state["cycle"], "tracks": len(new_track_ids), "mood": mood_data}

    except Exception as e:
        log.error(f"Cycle error: {e}", exc_info=True)
        return {"status": "error", "message": str(e)}
    finally:
        is_running = False
        STATE_LOCK.release()


def poll_recent_plays() -> bool:
    if is_running:
        return True
    with STATE_LOCK:
        try:
            state = load_state()
            sp = get_spotify()
            recent = sp.current_user_recently_played(limit=50)
        except Exception as e:
            log.error(f"Play counter poll failed: {e}")
            return False
        seen = set(state.get("seen_play_events", []))
        cum = dict(state.get("cumulative_plays", {}))
        new_events = 0
        for item in recent.get("items", []):
            tid = (item.get("track") or {}).get("id")
            if not tid:
                continue
            key = f'{item["played_at"]}|{tid}'
            if key in seen:
                continue
            seen.add(key)
            cum[tid] = cum.get(tid, 0) + 1
            new_events += 1
        if new_events:
            state["seen_play_events"] = sorted(seen)[-300:]
            state["cumulative_plays"] = cum
            save_state(state)
            log.info(f"Play counter: +{new_events} new plays (total {sum(cum.values())})")
        else:
            log.info("No new plays.")
        return True


def _hours_since_last_update(state: dict) -> float | None:
    last = state.get("last_update")
    if not last:
        return None
    try:
        last_dt = datetime.datetime.fromisoformat(last)
        if last_dt.tzinfo is None:
            last_dt = last_dt.replace(tzinfo=datetime.timezone.utc)
        now = datetime.datetime.now(datetime.timezone.utc)
        return (now - last_dt).total_seconds() / 3600
    except Exception:
        return None


if __name__ == "__main__":
    _mode = sys.argv[1].lower() if len(sys.argv) > 1 else ""
    if _mode == "poll":
        log.info("POLL mode.")
        if not poll_recent_plays():
            sys.exit(1)
    elif _mode == "cycle":
        log.info("CYCLE mode.")
        _hours = _hours_since_last_update(load_state())
        if _hours is not None:
            log.info(f"{_hours:.1f} hours since the last update.")
        result = run_cycle(manual=False)
        log.info(f"Cycle result: {result}")
        if result.get("status") != "ok":
            sys.exit(1)
    else:
        sys.exit("Usage: python bot.py [poll|cycle]")
